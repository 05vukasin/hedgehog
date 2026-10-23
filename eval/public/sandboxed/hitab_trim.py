"""Runs INSIDE the bubblewrap sandbox (venv-pandas: openpyxl 3.1.5). Writes trimmed copies of HiTab's
original annotation workbooks: only the table's sheet ('original' when present, else the first), with
every row from the first 'table descriptive sentence…' row in column A to the end removed (HiTab's
annotation block: the sentence, the question, the answer and its formula). Cached values are kept
(data_only=True), so formula cells become their values.

    python -I hitab_trim.py jobs.json results.jsonl      # jobs: [{"id", "src", "dst"}]; dst relative to /work
"""
import json
import sys
import warnings

warnings.filterwarnings("ignore")
import openpyxl  # noqa: E402

jobs = json.load(open(sys.argv[1]))
done = set()
try:
    done = {json.loads(l)["id"] for l in open(sys.argv[2])}
except FileNotFoundError:
    pass
out = open(sys.argv[2], "a")
for j in jobs:
    if j["id"] in done:
        continue
    r = {"id": j["id"]}
    try:
        wb = openpyxl.load_workbook(j["src"], data_only=True)
        name = "original" if "original" in wb.sheetnames else wb.sheetnames[0]
        for n in list(wb.sheetnames):
            if n != name:
                del wb[n]
        ws = wb[name]
        cut = None
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row, 1).value
            if isinstance(v, str) and v.strip().lower().startswith("table descriptive sentence"):
                cut = row
                break
        if cut:
            for m in list(ws.merged_cells.ranges):
                if m.max_row >= cut:
                    ws.merged_cells.remove(m)
            ws.delete_rows(cut, ws.max_row - cut + 1)
        import os
        os.makedirs(os.path.dirname(j["dst"]) or ".", exist_ok=True)
        wb.save(j["dst"])
        r |= {"ok": True, "sheet": name, "cut_row": cut, "rows": ws.max_row}
    except Exception as e:  # noqa: BLE001
        r |= {"ok": False, "error": f"{type(e).__name__}: {e}"[:300]}
    out.write(json.dumps(r) + "\n")
    out.flush()
