"""Runs INSIDE the bubblewrap sandbox. Naive baselines over untrusted spreadsheets, one forked child per
file (time limit, crash isolation), results in /work/<key>.json.gz.

    python -I baseline_dump.py pandas     /work/jobs.json   # venv-pandas: pandas 2.2.3 + openpyxl 3.1.5
    python -I baseline_dump.py markitdown /work/jobs.json   # venv-markitdown: markitdown 0.1.3

pandas: what a plain pandas tool gives an agent: pd.read_excel(path, sheet_name=None, header=0) (read_csv
for .csv), every sheet's column names and cells. `src_rows` maps each DataFrame row to its 1-based sheet row
(the sheet's first row is the header and blank rows are kept; two other layouts are tried), or null when
no layout gives the DataFrame's length.
markitdown: MarkItDown().convert(path).text_content and its length (tokens are counted outside, by score.py).
Peak RSS is the child's (includes the imported library, as a real tool process would).
"""
import gzip
import json
import math
import os
import resource
import sys
import time
import warnings

warnings.filterwarnings("ignore")
MODE, JOBS = sys.argv[1], sys.argv[2]
TIMEOUT = 300

if MODE == "pandas":
    import openpyxl
    import pandas as pd
else:
    from markitdown import MarkItDown
    MD = MarkItDown()


def cell(v):
    if v is None:
        return None
    if isinstance(v, float):
        return None if math.isnan(v) else (v if math.isfinite(v) else str(v))
    if isinstance(v, (int, str, bool)):
        return v
    if hasattr(v, "isoformat"):
        try:
            return v.isoformat()
        except Exception:  # noqa: BLE001  (NaT)
            return None
    return str(v)


def nonblank_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        first = 1  # read-only iter_rows starts at row 1 whatever the dimension says
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True), first):
            if any(v is not None and (not isinstance(v, str) or v != "") for v in row):
                rows.append(i)
        out[ws.title] = (first, rows)
    wb.close()
    return out


def do_pandas(path):
    t = time.perf_counter()
    if path.lower().endswith(".csv"):
        dfs = {"(csv)": pd.read_csv(path, header=0)}
    else:
        dfs = pd.read_excel(path, sheet_name=None, header=0)
    secs = time.perf_counter() - t
    try:
        nb = nonblank_rows(path) if path.lower().endswith((".xlsx", ".xlsm")) else {}
    except Exception:  # noqa: BLE001
        nb = {}
    sheets = []
    for name, df in dfs.items():
        first, rows = nb.get(name, (1, []))
        src = None
        if rows and rows[-1] - first == len(df):       # header = the sheet's first row, blank rows kept
            src = list(range(first + 1, rows[-1] + 1))
        elif rows and len(rows) - 1 == len(df):        # header = first non-blank row
            src = rows[1:]
        elif len([r for r in rows if r > first]) == len(df):  # header = the sheet's first row (blank)
            src = [r for r in rows if r > first]
        sheets.append({"name": name, "columns": [str(c) for c in df.columns], "src_rows": src,
                       "dtypes": [str(d) for d in df.dtypes],
                       "rows": [[cell(v) for v in r] for r in df.itertuples(index=False, name=None)]})
    return {"secs": round(secs, 4), "sheets": sheets}


def do_markitdown(path):
    t = time.perf_counter()
    text = MD.convert(path).text_content
    secs = time.perf_counter() - t
    return {"secs": round(secs, 4), "chars": len(text), "text": text}


def child(job):
    out = f"/work/{job['key']}.json.gz"
    try:
        r = (do_pandas if MODE == "pandas" else do_markitdown)(job["path"])
        r["error"] = None
    except Exception as e:  # noqa: BLE001
        r = {"error": f"{type(e).__name__}: {e}"[:400]}
    r["maxrss_kb"] = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
    with gzip.open(out + ".part", "wt") as f:
        json.dump(r, f, ensure_ascii=False)
    os.replace(out + ".part", out)


def main():
    jobs = json.load(open(JOBS))
    for j in jobs:
        out = f"/work/{j['key']}.json.gz"
        if os.path.exists(out):
            continue
        pid = os.fork()
        if pid == 0:
            try:
                child(j)
            finally:
                os._exit(0)
        t = time.time()
        status = None
        while time.time() - t < TIMEOUT:
            done, status = os.waitpid(pid, os.WNOHANG)
            if done:
                break
            time.sleep(0.02)
        else:
            os.kill(pid, 9)
            os.waitpid(pid, 0)
            status = "timeout"
        if not os.path.exists(out):
            err = "timeout after 300 s" if status == "timeout" else f"child died (status {status})"
            with gzip.open(out, "wt") as f:
                json.dump({"error": err}, f)


main()
