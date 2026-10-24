"""Cell-fidelity gate: is every source cell in hedgehog's normalized output, with the same value?

    PY=../scratch-archive/fidelity/venv/bin/python        # eval/fidelity/setup.sh makes it
    $PY -I eval/fidelity/check.py run --corpus mimotable [--bin target-master/release/hedgehog] [--workers 3]
    $PY -I eval/fidelity/check.py summary [--corpus a,b] [--out results.md]
    $PY -I eval/fidelity/check.py gate --baseline eval/fidelity/baseline.tsv [--update]
    $PY -I eval/fidelity/check.py quick --bin ... --tmp DIR [--update]      # make quick: 27 files
    $PY -I eval/fidelity/check.py full --bin ... [--update]                  # make fidelity: all + gate

Model-free and deterministic: no LLM, no network. Run it before any LLM evaluation: an LLM run
should only measure reasoning, never a misread date, a lost value or a changed ID.

Per file, every non-empty source cell is read independently of hedgehog (source.py: openpyxl /
xlrd / odfpy / csv, with the cell's number format) and looked up in hedgehog's output of
`hedgehog normalize FILE -o DIR --csv --keep-totals` through the manifest's cell mapping:
- table data: source cell = <columns[c].source_column><_src_row> on the table's source sheet;
  the value is compared (compare.py): numbers equal (integers exactly, floats to 1e-9 relative;
  percent cells as fraction or x100, reported), dates the same day and second (Excel serials left
  as numbers, day/month swaps, invented years), text exactly after whitespace collapsing (hedgehog
  documents trimming and collapsing), IDs and leading zeros kept, no scientific notation;
- header cells (columns[].original_header), titles, notes (the normalized workbook's Notes sheet,
  whose text is compared too), totals (its Totals sheet), key/value tables (`cell` column), group
  labels (the `group` column) and header words: placed;
- a missing-value mark ("x", "NA", "None", "(X)", "..") inside a table's source range is null there
  (a row or column of nothing but such marks is removed like an empty one): intentionally_dropped;
- a text cell whose value is the label split over it and the rows above/below joined: ok (B72);
- a flag column (`X` / `✓` marks and blanks in a table's range): each mark must be kept as a
  non-empty value, else text_changed (flag_lost);
- cells only on Notes as `other` that form a block of values (two or more rows with two numbers
  or dates): MISPLACED (value_block_in_notes), a table left out;
- anything else is UNPLACED; header cells that went to Notes as `repeated_header` of a table that
  got no header at all are MISPLACED (the table's columns are col_<letter>).
Column checks on top: collapsed dates (many distinct source dates, few distinct values) and a
day/month order that contradicts a value of the same column that only one order can read.

Classes per cell: ok, value_changed, date_changed, text_changed, unplaced, misplaced,
intentionally_dropped (missing tokens such as "-" / "n/a", Excel errors: documented).
Rows below --row-limit (10,000) per sheet are not checked (big registers).

Results: <work>/<corpus>/results.jsonl (one line per file, resumable), summary.json and SUMMARY.md
from `summary`. `gate` compares the corpus totals with baseline.tsv: a failure count may only go
down (like eval/quick/info.tsv); --update writes the current counts.
"""

from __future__ import annotations

import argparse
import warnings
import csv
import datetime as dt
import glob
import hashlib
import json
import multiprocessing as mp
import os
import re
import shutil
import subprocess
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from compare import compare, date_readings, is_missing_token, norm_loose, parse_hh_date, same_moment, show, ws  # noqa: E402
from source import Src, read_source  # noqa: E402

REPO = HERE.parent.parent
PROJECT = REPO.parent
csv.field_size_limit(1 << 30)

FAIL = ("value_changed", "date_changed", "text_changed", "unplaced", "misplaced")
CLASSES = ("ok",) + FAIL + ("intentionally_dropped",)
EXAMPLES_PER_SUB = 4
TIMEOUT = 300


def data_dir() -> Path:
    return Path(os.environ.get("HEDGEHOG_DATA") or PROJECT / "testdata")


def work_dir() -> Path:
    return Path(os.environ.get("HH_FIDELITY_WORK") or PROJECT / "scratch-archive" / "fidelity" / "work")


# ------------------------------------------------------------------------------------------------
# corpora: name -> [relative path under the data folder]

SPREADSHEET = (".xlsx", ".xlsm", ".xls", ".ods", ".csv")


def corpus_files(name: str) -> list[str]:
    d = data_dir()
    if name in ("mimotable", "realhitbench", "hitab"):
        sys.path.insert(0, str(REPO / "eval" / "public"))
        os.environ.setdefault("HEDGEHOG_DATA", str(d))
        from run_hedgehog import dataset_files  # noqa: E402
        return [f"public/{name}/{f}" for f in dataset_files(name)]
    if name == "tablesense":
        out = set()
        p = d / "public" / "tablesense" / "tables.tsv"
        for row in csv.DictReader(p.open(), delimiter="\t"):
            f = os.path.normpath(os.path.join("public/tablesense", row["file"]))
            out.add(f)
        return sorted(f for f in out if (d / f).exists())
    if name == "agents":
        cases = json.loads((REPO / "eval" / "agents" / "cases.json").read_text())
        out = set()
        for c in cases:
            root = c.get("root") or "irregular"
            for f in c.get("files") or [c["file"]]:
                out.add(f"{root}/{f}")
        return sorted(f for f in out if (d / f).exists())
    if name == "golden":
        out = []
        for lst in (REPO / "core" / "benches" / "enron50.txt", REPO / "eval" / "quick" / "files.txt"):
            for line in lst.read_text().splitlines():
                line = line.split("#")[0].strip()
                if line and (d / line).exists() and line not in out:
                    out.append(line)
        return out
    if name == "quick":
        return [ln.split("#")[0].strip() for ln in (HERE / "quick_files.txt").read_text().splitlines()
                if ln.split("#")[0].strip()]
    raise SystemExit(f"unknown corpus {name}")


CORPORA = ("mimotable", "realhitbench", "hitab", "tablesense", "agents", "golden")


def key_of(rel: str) -> str:
    return hashlib.sha1(rel.encode()).hexdigest()[:16]


# ------------------------------------------------------------------------------------------------
# hedgehog's output


def col_index(letter: str) -> int:
    n = 0
    for ch in letter:
        n = n * 26 + ord(ch.upper()) - 64
    return n


_A1 = re.compile(r"^([A-Z]+)(\d+)$")


def parse_a1(a: str):
    m = _A1.match(a.strip().replace("$", ""))
    return (int(m.group(2)), col_index(m.group(1))) if m else None


def a1(row: int, col: int) -> str:
    s = ""
    while col:
        col, r = divmod(col - 1, 26)
        s = chr(65 + r) + s
    return f"{s}{row}"


def run_normalize(binary: str, src: Path, out: Path) -> dict:
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    t = time.perf_counter()
    try:
        p = subprocess.run([binary, "normalize", str(src), "-o", str(out), "--csv", "--keep-totals"],
                           capture_output=True, text=True, timeout=TIMEOUT, env=dict(os.environ, HEDGEHOG_CACHE="0"))
        rc, err = p.returncode, p.stderr.strip()[-500:]
    except subprocess.TimeoutExpired:
        rc, err = -9, "timeout"
    return {"rc": rc, "err": err, "secs": round(time.perf_counter() - t, 3)}


def read_sheet_rows(xlsx: Path, names: list[str]) -> dict:
    """Rows (as dicts by header) of the normalized workbook's Notes / Totals sheets."""
    import openpyxl
    out = {}
    try:
        wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    except Exception:
        return out
    for n in names:
        if n not in wb.sheetnames:
            continue
        rows = wb[n].iter_rows(values_only=True)
        head = next(rows, None)
        if not head:
            continue
        out[n] = [dict(zip(head, r)) for r in rows]
    wb.close()
    return out


class HH:
    """Placements of hedgehog's output: (sheet, row, col) -> list of (role, info)."""

    def __init__(self, out: Path, row_limit: int):
        mpath = next(iter(glob.glob(str(out / "*.manifest.json"))), None)
        self.ok = mpath is not None
        self.place: dict = defaultdict(list)
        self.tables: list = []
        self.labels: dict = defaultdict(set)  # sheet -> normalized label words of groups / headers / titles
        self.headerless: dict = {}
        self.last_header: dict = {}
        self.ranges: dict = defaultdict(list)
        if not self.ok:
            return
        m = json.loads(Path(mpath).read_text())
        self.manifest = m
        csvs = {Path(p).name: p for p in glob.glob(str(out / "*.csv"))}
        for t in m.get("tables", []):
            s = t["source"]
            sheet = s["sheet"]
            self.tables.append(t)
            cols = t["columns"]
            title = s.get("title") or ""
            if title:
                self.labels[sheet].add(norm_loose(title))
            if s.get("title_cell"):
                rc = parse_a1(s["title_cell"])
                if rc:
                    self.place[(sheet, *rc)].append(("title", {"table": t["id"], "text": title}))
            for n in s.get("notes") or []:
                rc = parse_a1(n.get("cell") or "")
                if rc:
                    self.place[(sheet, *rc)].append(("note", {"table": t["id"], "text": n.get("text")}))
            kv = title == "key/value cells" and cols and cols[0]["name"] == "key"
            rollover = any("read as the next year" in w for w in t.get("warnings") or [])
            no_header = not s.get("header_range") and all(c["name"].startswith("col_") for c in cols if c.get("source_column"))
            if no_header:
                self.headerless[(sheet, t["id"])] = t
            for c in cols:
                for part in c["name"].split(" / "):
                    self.labels[sheet].add(norm_loose(part))
                for h in c.get("original_header") or []:
                    rc = parse_a1(h.get("cell") or "")
                    if rc:
                        self.place[(sheet, *rc)].append(("header", {"table": t["id"], "text": h.get("text"), "col": c["name"],
                                                                     "dtype": c.get("dtype"), "row": rc[0], "rollover": rollover}))
                        self.last_header[(sheet, t["id"])] = max(self.last_header.get((sheet, t["id"]), 0), rc[0])
            fname = next((n for n in csvs if n.endswith("." + t["sheet"] + ".csv")), None)
            if not fname:
                continue
            with open(csvs[fname], newline="", encoding="utf-8") as fh:
                rd = csv.reader(fh)
                head = next(rd, [])
                try:
                    isrc = head.index(t.get("source_row_column") or "_src_row")
                except ValueError:
                    continue
                srccols = [(i, col_index(c["source_column"]), c) for i, c in enumerate(cols)
                           if c.get("source_column") and i < len(head)]
                # Group columns, and the outline columns of group and row labels (`group_parent`, `<label>_parent`):
                # labels written once in the sheet.
                labcols = [i for i, c in enumerate(cols) if not c.get("source_column") and not c["name"].startswith("_")
                           and (not c.get("derived") or (c.get("derived") or {}).get("kind") == "parent") and i < len(head)]
                if kv:
                    ik = next((i for i, c in enumerate(cols) if c["name"] == "key"), None)
                    iv = next((i for i, c in enumerate(cols) if c["name"] == "value"), None)
                    ic = next((i for i, c in enumerate(cols) if c["name"] == "cell"), None)
                for row in rd:
                    try:
                        r = int(row[isrc])
                    except (ValueError, IndexError):
                        continue
                    if r > row_limit:
                        continue
                    if kv and None not in (ik, iv, ic):
                        rc = parse_a1(row[ic])
                        if rc:
                            self.place[(sheet, *rc)].append(("kv_value", {"table": t["id"], "value": row[iv], "dtype": cols[iv]["dtype"], "colmeta": cols[iv]}))
                            self.place[("#kvkey", sheet, rc[0])].append(row[ik])
                        continue
                    for i, cc, c in srccols:
                        self.place[(sheet, r, cc)].append(("data", {"table": t["id"], "value": row[i] if i < len(row) else "",
                                                                     "dtype": c["dtype"], "col": c["name"], "colmeta": c}))
                    for i in labcols:
                        if i < len(row) and row[i]:
                            self.labels[sheet].add(norm_loose(row[i]))  # one label (a group level of its own)
                            for part in row[i].split(" / "):
                                self.labels[sheet].add(norm_loose(part))
        # sheet -> [(top, left, bottom, right)] of the tables' source ranges
        for t in self.tables:
            a, _, b = (t["source"].get("range") or "").partition(":")
            ra, rb = parse_a1(a) if a else None, parse_a1(b or a) if a else None
            if ra and rb:
                self.ranges[t["source"]["sheet"]].append((ra[0], ra[1], rb[0], rb[1]))
        x = next(iter(glob.glob(str(out / "*.normalized.xlsx"))), None)
        extra = read_sheet_rows(Path(x), ["Notes", "Totals"]) if x else {}
        for n in extra.get("Notes", []):
            rc = parse_a1(str(n.get("cell") or ""))
            if rc and rc[0] <= row_limit:
                self.place[(n.get("sheet"), *rc)].append(("note", {"kind": n.get("kind"), "table": n.get("table"),
                                                                    "text": "" if n.get("text") is None else str(n.get("text"))}))
        for tt in extra.get("Totals", []):
            cell = str(tt.get("cell") or "")
            a, _, b = cell.partition(":")
            ra, rb = parse_a1(a), parse_a1(b) if b else None
            if not ra:
                continue
            for r in range(ra[0], min((rb or ra)[0], row_limit) + 1):
                self.place[(tt.get("sheet"), r, ra[1])].append(("total", {"table": tt.get("table")}))
            if tt.get("label_cell"):
                rc = parse_a1(str(tt["label_cell"]))
                if rc:
                    self.place[(tt.get("sheet"), *rc)].append(("total_label", {}))


    def in_table(self, sheet, r, c) -> bool:
        return any(t <= r <= b and l <= c <= rt for t, l, b, rt in self.ranges.get(sheet, ()))


# ------------------------------------------------------------------------------------------------
# one file


def check_file(binary: str, rel: str, src_path: Path, out: Path, row_limit: int) -> dict:
    res: dict = {"file": rel, "counts": Counter(), "subs": Counter(), "examples": defaultdict(list), "columns": []}
    norm = out / "norm"
    res["normalize"] = run_normalize(binary, src_path, norm)
    t0 = time.perf_counter()
    try:
        cells = read_source(src_path, row_limit)
    except Exception as e:  # an unreadable file is reported, not a failure of hedgehog
        res["source_error"] = f"{type(e).__name__}: {str(e)[:200]}"
        cells = None
    res["source_secs"] = round(time.perf_counter() - t0, 3)
    if res["normalize"]["rc"] != 0:
        res["hh_error"] = res["normalize"]["err"]
    if cells is None or res["normalize"]["rc"] != 0:
        return _finish(res, norm)
    hh = HH(norm, row_limit)
    if not hh.ok:
        res["hh_error"] = "no manifest"
        return _finish(res, norm)
    # CSV sources have one sheet, named by hedgehog
    if src_path.suffix.lower() in (".csv", ".tsv", ".txt"):
        names = {t["source"]["sheet"] for t in hh.tables} or {k[0] for k in hh.place if isinstance(k[0], str) and len(k) == 3}
        name = next(iter(sorted(names)), "")
        cells = {(name, r, c): v for (_, r, c), v in cells.items()}
    flags = flag_cells(cells, hh)
    other_notes: list = []
    pct = Counter()
    hdr_rows: dict = defaultdict(list)  # (sheet, table, row) -> header cells of a table's last header row
    colvals: dict = defaultdict(list)  # (table, col) -> [(src, hh, addr)]
    for (sheet, r, c), s in sorted(cells.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2])):
        addr = f"{sheet}!{a1(r, c)}"
        pl = hh.place.get((sheet, r, c), [])
        data = [p for p in pl if p[0] in ("data", "kv_value")]
        if data:
            # several tables may hold one cell (copied labels): the first that matches wins
            best = None
            for role, info in data:
                cls, sub = compare(s, info["value"], info["dtype"], info.get("colmeta"))
                if best is None or (best[0] != "ok" and cls == "ok"):
                    best = (cls, sub, info)
            cls, sub, info = best
            if (sheet, r, c) in flags:
                cls, sub = ("ok", "flag") if ws(info["value"]) else ("text_changed", "flag_lost")
            if (cls, sub) == ("text_changed", "lost") and norm_loose(s.value) in hh.labels.get(sheet, ()):
                cls, sub = "ok", "in_group_column"  # an inline group label: in the group column
            if cls == "text_changed" and s.kind == "text" and joined_label(cells, sheet, r, c, info["value"]):
                cls, sub = "ok", "label_joined"  # a label split over rows, joined on each (B72)
            if sub == "percent_fraction" or ("%" in (s.fmt or "") and s.kind == "number"):
                pct["fraction" if cls == "ok" else "x100" if sub == "percent_x100" else "other"] += 1
            colvals[(info["table"], info["col"] if "col" in info else "value")].append((s, info["value"], addr, info["dtype"]))
            _count(res, cls, sub, addr, s, info["value"], f'{info["table"]}.{info.get("col", "value")}')
            continue
        notes = [p for p in pl if p[0] == "note"]
        heads = [p for p in pl if p[0] == "header"]
        if heads:
            h = heads[0][1]
            cls, sub = _cmp_header(s, h)
            if cls == "ok" and h["row"] == hh.last_header.get((sheet, h["table"])) and h["row"] == r:
                hdr_rows[(sheet, h["table"], r)].append((addr, s, h))
            _count(res, cls, sub, addr, s, h.get("text"), f'{h["table"]} header')
            continue
        if notes:
            n = notes[0][1]
            if n.get("kind") == "repeated_header" and any((sheet, n.get("table")) == k for k in hh.headerless):
                _count(res, "misplaced", "header_demoted_to_notes", addr, s, n.get("text"), f'{n.get("table")} has no header')
                continue
            cls, sub = _cmp_note(s, n.get("text") or "")
            if (sheet, r, c) in flags:
                _count(res, "text_changed", "flag_lost", addr, s, n.get("text"), f'Notes ({n.get("kind")})')
                continue
            _count(res, cls, "note:" + sub if cls != "ok" else "note", addr, s, n.get("text"), f'Notes ({n.get("kind")})')
            if cls == "ok" and n.get("kind") == "other":
                other_notes.append((sheet, r, c, s, addr))
            continue
        if pl:
            _count(res, "ok", pl[0][0], addr, s, "", pl[0][0])
            continue
        # labels: a key of a key/value table, a group label, a word of a column name or title
        if s.kind == "text":
            k = norm_loose(s.value)
            kv_keys = hh.place.get(("#kvkey", sheet, r), [])
            if any(k and k in norm_loose(x) for x in kv_keys):
                _count(res, "ok", "kv_key", addr, s, "", "key/value")
                continue
            if k and k in hh.labels.get(sheet, ()):
                _count(res, "ok", "label", addr, s, "", "group/header label")
                continue
            # "Region: North" group labels: the column is named after the prefix (documented)
            pre, sep, val = s.value.partition(":")
            if sep and norm_loose(val) in hh.labels.get(sheet, ()) and norm_loose(pre) in hh.labels.get(sheet, ()):
                _count(res, "ok", "label_prefix", addr, s, "", "group label")
                continue
            if not k:
                _count(res, "intentionally_dropped", "punctuation_only", addr, s, "", "")
                continue
        # a whole number used as a section label (a year row `2015` over its rows): in the group column
        if s.kind == "number" and isinstance(s.value, (int, float)) and float(s.value).is_integer() \
                and norm_loose(str(int(s.value))) in hh.labels.get(sheet, ()):
            _count(res, "ok", "label", addr, s, "", "group label (number)")
            continue
        if s.kind == "error":
            _count(res, "intentionally_dropped", "excel_error_outside_tables", addr, s, "", "")
            continue
        if (sheet, r, c) in flags:
            _count(res, "text_changed", "flag_lost", addr, s, "", "flag column")
            continue
        # A missing-value mark ("x", "NA", "None", "..") inside a table's range: it is null there, and a
        # row or column of nothing but such marks is removed like an empty one (documented drop).
        if s.kind == "text" and (is_missing_token(s.value) or s.value.strip().lower() == "none") and hh.in_table(sheet, r, c):
            _count(res, "intentionally_dropped", "missing_token_in_table", addr, s, "", "")
            continue
        if s.kind in ("date", "datetime") and norm_loose(show(s.value)) in hh.labels.get(sheet, ()):
            _count(res, "ok", "label_date", addr, s, "", "header label")  # a repeated date header row
            continue
        _count(res, "unplaced", s.kind, addr, s, "", "")
    _header_data_rows(res, hdr_rows)
    _value_blocks_in_notes(res, other_notes)
    _column_checks(res, colvals)
    res["percent"] = dict(pct)
    return _finish(res, norm)


YEARLIKE = range(1800, 2101)


FLAG_MARKS = {"x", "✓", "✔"}


def numeric_format(fmt) -> bool:
    """A number format with digit placeholders (`#,##0.0`, `0.000`): the cell stands for a number,
    so an `X` in it is a missing-value mark (Destatis "Tabellenfach gesperrt"), not a flag."""
    f = (fmt or "").strip()
    return bool(f) and f.lower() not in ("general", "@") and bool(re.search(r"[0#]", re.sub(r'"[^"]*"', "", f)))


def flag_cells(cells: dict, hh) -> set:
    """Cells of flag columns: in a table's source range, a column whose filled cells are `x` / `X`
    / `✓` marks (at least two; one other text on top, its header, allowed) with blank rows
    between. The mark is a value (flagged vs not), never a missing-value mark: it must be kept,
    distinct from blank."""
    out = set()
    marked_cols = defaultdict(set)  # sheet -> columns holding a mark at all
    by_row = defaultdict(lambda: defaultdict(list))  # sheet -> row -> columns
    for (sh, r, c), x in cells.items():
        by_row[sh][r].append(c)
        if x.kind == "text" and x.value.strip().lower() in FLAG_MARKS:
            marked_cols[sh].add(c)
    for sheet, ranges in hh.ranges.items():
        if not marked_cols.get(sheet):
            continue
        for top, left, bottom, right in ranges:
            cols = [c for c in marked_cols[sheet] if left <= c <= right]
            if not cols:
                continue
            # rows of the table with two or more cells (not a lone group label or caption)
            rows_with = sorted(r for r, cs in by_row[sheet].items() if top <= r <= bottom and sum(left <= c <= right for c in cs) >= 2)
            for col in cols:
                marks, others = [], []
                for r in rows_with:
                    x = cells.get((sheet, r, col))
                    if x is None:
                        continue
                    if x.kind == "text" and x.value.strip().lower() in FLAG_MARKS and not numeric_format(x.fmt):
                        marks.append(r)
                    else:
                        others.append(r)
                # blank rows among the marked ones (rows above the first mark are its header)
                blanks = sum(1 for r in rows_with if marks and r > marks[0]) - sum(1 for r in marks[1:]) \
                    - sum(1 for r in others if marks and r > marks[0])
                header_only = not others or (len(others) == 1 and marks and others[0] < marks[0])
                if len(marks) >= 2 and header_only and blanks >= 1:
                    out.update((sheet, r, col) for r in marks)
    return out


def _value_blocks_in_notes(res, other_notes):
    """Cells placed only as Notes items of kind `other` that form a block of values: row segments
    (cells at most two blank columns apart) joined over consecutive rows with overlapping columns,
    with at least two rows holding two or more numbers or dates in one segment: a table that
    hedgehog did not make a table. MISPLACED (value_block_in_notes)."""
    by_sheet = defaultdict(lambda: defaultdict(list))
    for sheet, r, c, s, addr in other_notes:
        by_sheet[sheet][r].append((c, s, addr))
    for sheet, rows in by_sheet.items():
        segs = []  # (row, left, right, cells)
        for r in sorted(rows):
            for c, s, addr in sorted(rows[r], key=lambda x: x[0]):
                if segs and segs[-1][0] == r and c <= segs[-1][2] + 2:
                    segs[-1][2] = c
                    segs[-1][3].append((c, s, addr))
                else:
                    segs.append([r, c, c, [(c, s, addr)]])
        parent = list(range(len(segs)))

        def find(x):
            while parent[x] != x:
                parent[x] = parent[parent[x]]
                x = parent[x]
            return x
        by_row = defaultdict(list)
        for k, sg in enumerate(segs):
            by_row[sg[0]].append(k)
        for k, sg in enumerate(segs):
            for j in by_row.get(sg[0] - 1, ()):
                if segs[j][1] <= sg[2] + 1 and sg[1] <= segs[j][2] + 1:
                    parent[find(k)] = find(j)
        comps = defaultdict(list)
        for k in range(len(segs)):
            comps[find(k)].append(segs[k])
        for comp in comps.values():
            value_rows = sum(1 for sg in comp if sum(1 for _, s, _ in sg[3] if s.kind in ("number", "date", "datetime")) >= 2)
            if len({sg[0] for sg in comp}) < 2 or value_rows < 2:
                continue
            for sg in comp:
                for c, s, addr in sg[3]:
                    res["counts"]["ok"] -= 1
                    res["counts"]["misplaced"] = res["counts"].get("misplaced", 0) + 1
                    res["subs"]["misplaced:value_block_in_notes"] += 1
                    ex = res["examples"]["misplaced:value_block_in_notes"]
                    if len(ex) < EXAMPLES_PER_SUB:
                        ex.append({"cell": addr, "source": _show(s), "hedgehog": "", "where": "Notes (other)"})


def joined_label(cells: dict, sheet: str, r: int, c: int, hv) -> bool:
    """hedgehog's value is this text cell joined with the text cells right above and/or below it
    in the same column: a row label split over two to five rows ("CTFP" / "Por tempo" /
    "indeterminado"), written in full on each of its rows (documented, B72)."""
    want = ws(hv or "")
    if not want or len(want) > 400:
        return False
    for top in range(r, r - 5, -1):
        for bot in range(r, top + 5):
            if bot == top:
                continue
            parts = []
            for rr in range(top, bot + 1):
                x = cells.get((sheet, rr, c))
                if x is None or x.kind != "text":
                    break
                parts.append(ws(x.value))
            else:
                if " ".join(parts) == want:
                    return True
    return False


def _header_data_rows(res, hdr_rows):
    """A table's last header row that reads like a data row (a date over a date column, a number
    that is no year over a number column, and text over text columns: at least three such cells)
    lost its values into the column names: those cells are MISPLACED (data_row_in_header)."""
    numeric = ("integer", "number", "currency", "percent")
    for (sheet, tid, row), cells in hdr_rows.items():
        match = strong = 0
        for addr, s, h in cells:
            dt_ = h.get("dtype") or ""
            if s.kind in ("date", "datetime") and dt_ in ("date", "datetime"):
                match += 1
                strong += 1
            elif s.kind == "number" and dt_ in numeric:
                match += 1
                if not (float(s.value).is_integer() and int(s.value) in YEARLIKE):
                    strong += 1
            elif s.kind == "text" and dt_ == "text":
                match += 1
        nums = [float(s.value) for _, s, _ in cells if s.kind == "number"]  # cells come in column order
        ints = [int(x) for x in nums if x.is_integer()]
        numbering = len(ints) >= 3 and len(ints) == len(nums) and all(b - a == 1 for a, b in zip(ints, ints[1:]))
        if match >= 3 and strong >= 1 and match * 2 >= len(cells) + 1 and not numbering:
            for addr, s, h in cells:
                res["counts"]["ok"] -= 1
                res["counts"]["misplaced"] = res["counts"].get("misplaced", 0) + 1
                res["subs"]["misplaced:data_row_in_header"] += 1
                ex = res["examples"]["misplaced:data_row_in_header"]
                if len(ex) < EXAMPLES_PER_SUB:
                    ex.append({"cell": addr, "source": _show(s), "hedgehog": h.get("col"), "where": f"{tid} header row {row}"})


def _cmp_header(s: Src, h: dict) -> tuple[str, str]:
    text = h.get("text") or ""
    if s.kind == "text":
        if ws(s.value) == ws(text) or norm_loose(s.value) == norm_loose(text):
            return "ok", "header"
        if norm_loose(text) and (norm_loose(text) in norm_loose(s.value) or norm_loose(s.value) in norm_loose(text)):
            return "ok", "header_cleaned"
        return "ok", "header_rewritten"
    if s.kind in ("date", "datetime", "time"):
        d = parse_hh_date(text)
        if d is not None and same_moment(s.value, d):
            return "ok", "header_date"
        rd, _ = date_readings(text)
        if any(same_moment(s.value, x) for x, _ in rd):
            return "ok", "header_date"
        if s.raw is not None and re.fullmatch(r"\d+(\.\d+)?", text.strip()):
            return "date_changed", "header:serial_left_as_number"
        if d is not None and isinstance(s.value, dt.date) and h.get("rollover"):
            a = s.value if not isinstance(s.value, dt.datetime) else s.value.date()
            b = d if not isinstance(d, dt.datetime) else d.date()
            if b.year == a.year + 1 and (b.month, b.day) == (a.month, a.day):
                return "ok", "header_year_rollover"  # documented (B84): typed without the year change
        if s.raw == 0 and text.strip() in ("0", "00:00:00"):
            return "ok", "zero_serial"
        return "date_changed", "header:date_mismatch"
    if s.kind == "number":
        cls, sub = compare(s, text.strip(), "text")
        return (cls, "header:" + sub) if cls != "ok" else ("ok", "header_number")
    return "ok", "header"


def _cmp_note(s: Src, text: str) -> tuple[str, str]:
    if s.kind == "text":
        if ws(s.value) == ws(text) or norm_loose(s.value) == norm_loose(text):
            return "ok", "note"
        return "text_changed", "text_mismatch"
    return compare(s, text, "text")


def _count(res, cls, sub, addr, s, hv, where):
    res["counts"][cls] += 1
    res["subs"][f"{cls}:{sub}"] += 1
    if cls not in ("ok", "intentionally_dropped") or sub in ("percent_fraction",):
        ex = res["examples"][f"{cls}:{sub}"]
        if len(ex) < EXAMPLES_PER_SUB:
            ex.append({"cell": addr, "source": _show(s), "hedgehog": hv, "where": where})


def _show(s: Src) -> str:
    v = show(s.value)
    return f"{v} [{s.kind}{', ' + s.fmt if s.fmt else ''}]"


def _column_checks(res, colvals):
    """Collapsed dates and day/month order per output column."""
    for (tid, col), vals in colvals.items():
        dates = [(s, hv, addr) for s, hv, addr, dtype in vals if dtype in ("date", "datetime")]
        if len(dates) < 3:
            continue
        src_keys, hh_keys = set(), set()
        for s, hv, addr in dates:
            if s.kind in ("date", "datetime"):
                src_keys.add(show(s.value))
            elif s.kind == "text":
                src_keys.add(ws(s.value))
            else:
                src_keys.add(show(s.value))
            if hv:
                hh_keys.add(hv)
        if len(src_keys) >= 4 and len(hh_keys) * 2 <= len(src_keys):
            res["columns"].append({"table": tid, "column": col, "issue": "date_collapsed", "source_distinct": len(src_keys),
                                   "hedgehog_distinct": len(hh_keys), "example": [dates[0][2], ws(str(dates[0][0].value)), dates[0][1]]})
        # order: text dates a/b/y; if some value is readable only day-first (or only month-first),
        # a value read the other way is a swap
        only_dm = only_md = 0
        for s, hv, addr in dates:
            if s.kind != "text":
                continue
            tags = {t for _, t in date_readings(s.value)[0]}
            if tags == {"dmy"}:
                only_dm += 1
            elif tags == {"mdy"}:
                only_md += 1
        want = "dmy" if only_dm and not only_md else "mdy" if only_md and not only_dm else None
        if want is None:
            continue
        for s, hv, addr in dates:
            if s.kind != "text":
                continue
            rd = date_readings(s.value)[0]
            if len(rd) < 2:
                continue
            d = parse_hh_date(hv)
            got = next((t for x, t in rd if d is not None and same_moment(x, d)), None)
            if got and got != want:
                res["counts"]["ok"] -= 1
                res["counts"]["date_changed"] += 1
                res["subs"][f"date_changed:day_month_swap_in_column"] += 1
                ex = res["examples"]["date_changed:day_month_swap_in_column"]
                if len(ex) < EXAMPLES_PER_SUB:
                    ex.append({"cell": addr, "source": ws(s.value), "hedgehog": hv, "where": f"{tid}.{col} reads {want}"})
                okk = [k for k in res["subs"] if k.startswith("ok:text_date_ambiguous")]
                for k in okk:
                    if res["subs"][k] > 0:
                        res["subs"][k] -= 1
                        break


def _finish(res, norm: Path) -> dict:
    res["counts"] = dict(res["counts"])
    res["subs"] = dict(res["subs"])
    res["examples"] = dict(res["examples"])
    return res


def _job(args):
    warnings.filterwarnings("ignore")  # openpyxl's notes on unsupported extensions
    binary, rel, row_limit = args
    src = data_dir() / rel
    out = work_dir() / "tmp" / key_of(rel)
    try:
        r = check_file(binary, rel, src, out, row_limit)
    except Exception as e:
        r = {"file": rel, "crash": f"{type(e).__name__}: {e}"[:400], "counts": {}, "subs": {}, "examples": {}}
    shutil.rmtree(out, ignore_errors=True)
    return r


# ------------------------------------------------------------------------------------------------
# commands


def default_bin() -> str:
    return os.environ.get("HEDGEHOG_BIN") or str(REPO / "target-master" / "release" / "hedgehog")


def cmd_run(a) -> int:
    binary = str(Path(a.bin or default_bin()).resolve())
    for corpus in a.corpus.split(","):
        files = corpus_files(corpus)
        if a.limit:
            files = files[: a.limit]
        if a.files:
            want = set(a.files.split(","))
            files = [f for f in files if f in want or Path(f).name in want]
        out = work_dir() / corpus
        out.mkdir(parents=True, exist_ok=True)
        rp = out / "results.jsonl"
        if a.fresh and rp.exists():
            rp.unlink()
        done = set()
        if rp.exists():
            for line in rp.read_text().splitlines():
                try:
                    done.add(json.loads(line)["file"])
                except (ValueError, KeyError):
                    pass
        todo = [f for f in files if f not in done]
        print(f"{corpus}: {len(files)} files, {len(todo)} to do, binary {binary}", flush=True)
        t = time.time()
        with rp.open("a") as fo, mp.get_context("spawn").Pool(a.workers, maxtasksperchild=25) as pool:
            for i, r in enumerate(pool.imap_unordered(_job, [(binary, f, a.row_limit) for f in todo]), 1):
                fo.write(json.dumps(r, ensure_ascii=False, default=str) + "\n")
                fo.flush()
                if i % 50 == 0 or i == len(todo):
                    print(f"  {i}/{len(todo)} {time.time() - t:.0f} s", flush=True)
    return 0


def load_results(corpus: str) -> list[dict]:
    rp = work_dir() / corpus / "results.jsonl"
    if not rp.exists():
        return []
    by = {}
    for line in rp.read_text().splitlines():
        try:
            r = json.loads(line)
        except ValueError:
            continue
        by[r["file"]] = r
    want = set(corpus_files(corpus))
    return [by[f] for f in sorted(by) if f in want]


def aggregate(results: list[dict]) -> dict:
    tot, subs = Counter(), Counter()
    ex = defaultdict(list)
    errors = Counter()
    pct = Counter()
    cols = []
    for r in results:
        if r.get("crash"):
            errors["checker_crash"] += 1
        if r.get("source_error"):
            errors["source_unreadable"] += 1
        if r.get("hh_error"):
            errors["hedgehog_error"] += 1
        tot.update(r.get("counts", {}))
        subs.update(r.get("subs", {}))
        pct.update(r.get("percent", {}))
        for k, v in r.get("examples", {}).items():
            seen_files = {e["file"] for e in ex[k]}
            for e in v:
                if len(ex[k]) < 3 and r["file"] not in seen_files:
                    ex[k].append(dict(e, file=r["file"]))
                    seen_files.add(r["file"])
        for c in r.get("columns", []):
            cols.append(dict(c, file=r["file"]))
    files_with = Counter()
    for r in results:
        for k in FAIL:
            if r.get("counts", {}).get(k):
                files_with[k] += 1
    return {"files": len(results), "cells": sum(tot.values()), "counts": {k: tot.get(k, 0) for k in CLASSES},
            "files_with": dict(files_with), "subs": dict(subs.most_common()), "examples": dict(ex),
            "errors": dict(errors), "percent": dict(pct), "columns": cols}


def cmd_summary(a) -> int:
    corpora = a.corpus.split(",") if a.corpus else [c for c in CORPORA if (work_dir() / c / "results.jsonl").exists()]
    summary = {c: aggregate(load_results(c)) for c in corpora}
    out = Path(a.out) if a.out else work_dir() / "SUMMARY.md"
    (out.with_suffix(".json")).write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=str))
    lines = ["# Cell fidelity", "", "| corpus | files | cells | " + " | ".join(CLASSES) + " | errors |",
             "|---|---:|---:|" + "---:|" * len(CLASSES) + "---|"]
    for c, s in summary.items():
        lines.append(f"| {c} | {s['files']} | {s['cells']} | " + " | ".join(str(s["counts"][k]) for k in CLASSES)
                     + f" | {', '.join(f'{k} {v}' for k, v in s['errors'].items()) or '-'} |")
    lines.append("")
    allsubs = Counter()
    allex = defaultdict(list)
    for c, s in summary.items():
        for k, v in s["subs"].items():
            if k.split(":")[0] in FAIL:
                allsubs[k] += v
                for e in s["examples"].get(k, []):
                    if len(allex[k]) < 3:
                        allex[k].append(e)
    lines += ["## Top failure classes", ""]
    for k, v in allsubs.most_common(a.top):
        lines.append(f"### {k}: {v} cells")
        for e in allex[k]:
            lines.append(f"- `{e['file']}` {e['cell']}: source `{e['source']}` → hedgehog `{e['hedgehog']}` ({e['where']})")
        lines.append("")
    cols = [dict(x, corpus=c) for c, s in summary.items() for x in s["columns"]]
    if cols:
        lines += ["## Column issues", ""]
        for x in cols[:30]:
            lines.append(f"- {x['corpus']} `{x['file']}` {x['table']}.{x['column']}: {x['issue']} "
                         f"({x.get('source_distinct')} distinct source dates → {x.get('hedgehog_distinct')}; e.g. {x.get('example')})")
        lines.append("")
    lines += ["## Percent cells", ""]
    for c, s in summary.items():
        if s["percent"]:
            lines.append(f"- {c}: {s['percent']} (hedgehog stores percent cells as fractions)")
    out.write_text("\n".join(lines) + "\n")
    print("\n".join(lines[: 6 + len(summary)]))
    print(f"written {out} and {out.with_suffix('.json')}")
    return 0


GATE_KEYS = ("files", "cells") + CLASSES + ("errors",)


def cmd_gate(a) -> int:
    base = Path(a.baseline)
    corpora = a.corpus.split(",") if a.corpus else list(CORPORA)
    cur = {}
    for c in corpora:
        res = load_results(c)
        if not res:
            continue
        s = aggregate(res)
        cur[c] = {"files": s["files"], "cells": s["cells"], **s["counts"], "errors": sum(s["errors"].values())}
    if a.update:
        old = _read_tsv(base)
        old.update(cur)
        _write_tsv(base, old)
        print(f"fidelity: baseline updated ({', '.join(cur)})")
        return 0
    old = _read_tsv(base)
    bad = []
    for c, v in cur.items():
        b = old.get(c)
        if not b:
            bad.append(f"{c}: not in the baseline (run with --update)")
            continue
        for k in FAIL + ("errors",):
            if v[k] > b[k]:
                bad.append(f"{c}: {k} {b[k]} -> {v[k]}")
        if v["files"] != b["files"]:
            bad.append(f"{c}: files {b['files']} -> {v['files']}")
    line = " · ".join(f"{c} " + "/".join(str(cur[c][k]) for k in FAIL) for c in cur)
    if bad:
        print(f"fidelity FAILED ({' / '.join(FAIL)}): {line}")
        for b in bad:
            print("  " + b)
        return 1
    better = [f"{c} {k} {old[c][k]}->{cur[c][k]}" for c in cur for k in FAIL if cur[c][k] < old.get(c, {}).get(k, 0)]
    print(f"fidelity OK ({' / '.join(FAIL)}): {line}" + (f"; improved: {', '.join(better)} (run --update)" if better else ""))
    return 0


def _read_tsv(p: Path) -> dict:
    out = {}
    if not p.exists():
        return out
    rows = [ln.split("\t") for ln in p.read_text().splitlines() if ln and not ln.startswith("#")]
    head = rows[0]
    for r in rows[1:]:
        out[r[0]] = {k: int(v) for k, v in zip(head[1:], r[1:])}
    return out


def _write_tsv(p: Path, d: dict, key="corpus"):
    keys = GATE_KEYS if key == "corpus" else QUICK_KEYS
    lines = ["\t".join((key,) + keys)]
    for c in sorted(d):
        lines.append("\t".join([c] + [str(d[c].get(k, 0)) for k in keys]))
    p.write_text("\n".join(lines) + "\n")


QUICK_KEYS = CLASSES


def cmd_quick(a) -> int:
    """make quick: the files of quick_files.txt, per-file counts vs quick.tsv (failures may only drop)."""
    os.environ["HH_FIDELITY_WORK"] = str(Path(a.tmp).resolve() / "fidelity")
    files = [f for f in corpus_files("quick") if (data_dir() / f).exists()]
    if not files:
        print("fidelity: skipped (no test data)")
        return 0
    binary = str(Path(a.bin).resolve())
    with mp.get_context("spawn").Pool(a.workers) as pool:
        results = pool.map(_job, [(binary, f, a.row_limit) for f in files])
    cur = {}
    for r in results:
        c = r.get("counts", {})
        cur[r["file"]] = {k: c.get(k, 0) for k in CLASSES}
        if r.get("crash") or r.get("hh_error") or r.get("source_error"):
            cur[r["file"]]["value_changed"] += 10**6  # an error fails the slice
    base = HERE / "quick.tsv"
    if a.update:
        _write_tsv(base, cur, key="file")
        print(f"fidelity: quick.tsv updated ({len(cur)} files)")
        return 0
    old = _read_tsv(base)
    bad = []
    for f, v in cur.items():
        b = old.get(f)
        if b is None:
            bad.append(f"{f}: not in quick.tsv (make quick-update)")
            continue
        for k in FAIL:
            if v[k] > b[k]:
                r = next(x for x in results if x["file"] == f)
                exs = [e for kk, es in r.get("examples", {}).items() if kk.startswith(k) for e in es][:2]
                bad.append(f"{f}: {k} {b[k]} -> {v[k]}" + "".join(f"\n    {e['cell']}: {e['source']} -> {e['hedgehog']!r}" for e in exs))
    tot = Counter()
    for v in cur.values():
        tot.update(v)
    line = f"{len(cur)} files, {sum(tot.values())} cells: " + ", ".join(f"{k} {tot[k]}" for k in CLASSES if tot[k])
    if bad:
        print(f"fidelity FAILED: {line}")
        for b in bad:
            print("  " + b)
        return 1
    better = [f for f in cur if f in old and any(cur[f][k] < old[f][k] for k in FAIL)]
    print(f"fidelity OK: {line}" + (f" (improved in {len(better)} files: make quick-update)" if better else ""))
    return 0


def cmd_full(a) -> int:
    """make fidelity: every corpus from scratch, the summary, then the gate against baseline.tsv."""
    for c in CORPORA:
        if not corpus_files(c):
            print(f"{c}: no files on disk, skipped")
            continue
        cmd_run(argparse.Namespace(corpus=c, bin=a.bin, workers=a.workers, limit=None, files=None, fresh=True,
                                   row_limit=a.row_limit))
    cmd_summary(argparse.Namespace(corpus=None, out=None, top=25))
    return cmd_gate(argparse.Namespace(baseline=a.baseline, corpus=None, update=a.update))


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("--corpus", required=True)
    r.add_argument("--bin")
    r.add_argument("--workers", type=int, default=3)
    r.add_argument("--limit", type=int)
    r.add_argument("--files")
    r.add_argument("--fresh", action="store_true")
    r.add_argument("--row-limit", type=int, default=10_000)
    s = sub.add_parser("summary")
    s.add_argument("--corpus")
    s.add_argument("--out")
    s.add_argument("--top", type=int, default=20)
    g = sub.add_parser("gate")
    g.add_argument("--baseline", default=str(HERE / "baseline.tsv"))
    g.add_argument("--corpus")
    g.add_argument("--update", action="store_true")
    q = sub.add_parser("quick")
    q.add_argument("--bin", required=True)
    q.add_argument("--data")
    q.add_argument("--tmp", required=True)
    q.add_argument("--workers", type=int, default=3)
    q.add_argument("--row-limit", type=int, default=2_000)
    q.add_argument("--update", action="store_true")
    f = sub.add_parser("full")
    f.add_argument("--bin")
    f.add_argument("--workers", type=int, default=3)
    f.add_argument("--row-limit", type=int, default=10_000)
    f.add_argument("--baseline", default=str(HERE / "baseline.tsv"))
    f.add_argument("--update", action="store_true")
    a = ap.parse_args()
    if getattr(a, "data", None):
        os.environ["HEDGEHOG_DATA"] = str(Path(a.data).resolve())
    return {"run": cmd_run, "summary": cmd_summary, "gate": cmd_gate, "quick": cmd_quick, "full": cmd_full}[a.cmd](a)


if __name__ == "__main__":
    sys.exit(main())
