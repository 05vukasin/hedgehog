"""Independent readers for the cell-fidelity gate: every non-empty source cell, read without hedgehog.

    cells = read_source(path, row_limit)   # {(sheet, row, col): Src}, 1-based rows and columns

xlsx / xlsm: openpyxl (read_only, data_only: cached formula values) plus each cell's number format;
xls: xlrd (with formatting_info when it works); ods: odfpy; csv / tsv / txt: the csv module.
Dates are decided from the number format here (`is_date_format`), not by the library, so a serial
shown through a format openpyxl does not know (`[$-804]aaa;@`, a Chinese weekday) still counts as
a date. Files are untrusted: they are only parsed.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass
from pathlib import Path

csv.field_size_limit(1 << 26)


@dataclass
class Src:
    kind: str  # number | date | datetime | time | text | bool | error
    value: object  # float/int, datetime.date/datetime, datetime.time, str, bool
    fmt: str = ""  # number format ("" for General / unknown)
    raw: object = None  # the serial for a date-formatted number


# ------------------------------------------------------------------------------------------------
# number formats

_QUOTED = re.compile(r'"[^"]*"|\\.|_.|\*.')
_BRACKET = re.compile(r"\[[^\]]*\]")


def is_date_format(fmt: str) -> bool:
    """True for a number format that shows a date or time: y, m, d, h, s outside quotes and
    brackets ([h] / [mm] / [ss] elapsed time counts), or the CJK weekday codes aaa / aaaa."""
    if not fmt or fmt.lower() in ("general", "@"):
        return False
    f = fmt.split(";")[0]
    elapsed = re.search(r"\[(h+|m+|s+)\]", f, re.I) is not None
    f = _BRACKET.sub("", _QUOTED.sub("", f)).lower()
    if elapsed:
        return True
    if "aaa" in f:
        return True
    if re.search(r"[0#?]", f) and not re.search(r"[ydhs]", f):
        return False  # numbers ("0.00", "#,##0") with an "e" or "m" in a literal stay numbers
    return re.search(r"[ydmhs]", f) is not None


def is_time_only_format(fmt: str) -> bool:
    if re.search(r"\[(h+|m+|s+)\]", fmt or "", re.I):
        return True  # elapsed time [h]:mm
    f = _BRACKET.sub("", _QUOTED.sub("", fmt.split(";")[0])).lower()
    return re.search(r"[hs]", f) is not None and re.search(r"[yd]", f) is None and "aaa" not in f


def is_percent_format(fmt: str) -> bool:
    return "%" in _QUOTED.sub("", fmt or "")


def serial_to_datetime(serial: float, epoch1904: bool = False):
    if epoch1904:
        base = dt.datetime(1904, 1, 1)
    elif serial < 60:
        base = dt.datetime(1899, 12, 31)
    else:
        base = dt.datetime(1899, 12, 30)
    try:
        return base + dt.timedelta(days=float(serial))
    except (OverflowError, ValueError):
        return None


def _date_cell(serial: float, fmt: str, epoch1904: bool) -> Src | None:
    if serial == 0:
        return Src("date", dt.date(1899, 12, 31), fmt, 0)
    if serial < 0:
        return Src("number", serial, fmt)  # Excel shows ####
    d = serial_to_datetime(serial, epoch1904)
    if d is None:
        return None
    d = _round_second(d)
    if is_time_only_format(fmt) and serial < 1:
        return Src("time", d.time(), fmt, serial)
    if d.time() == dt.time(0):
        return Src("date", d.date(), fmt, serial)
    return Src("datetime", d, fmt, serial)


def _round_second(d: dt.datetime) -> dt.datetime:
    if d.microsecond >= 500_000:
        d = d + dt.timedelta(seconds=1)
    return d.replace(microsecond=0)


# ------------------------------------------------------------------------------------------------
# readers


def read_source(path: Path, row_limit: int = 10_000) -> dict:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path, row_limit)
    if ext == ".xls":
        try:
            return read_xls(path, row_limit)
        except Exception as e:  # an .xls that is really xlsx / html / csv
            head = path.read_bytes()[:8]
            if head.startswith(b"PK"):
                return read_xlsx(path, row_limit)
            raise e
    if ext == ".ods":
        return read_ods(path, row_limit)
    if ext in (".csv", ".tsv", ".txt"):
        return read_csv(path, row_limit)
    raise ValueError(f"unsupported extension {ext}")


def read_xlsx(path: Path, row_limit: int) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    epoch1904 = getattr(wb, "epoch", None) is not None and wb.epoch.year == 1904
    out: dict = {}
    for ws in wb.worksheets:
        if not hasattr(ws, "iter_rows"):
            continue  # chartsheet
        try:
            ws.reset_dimensions()
        except AttributeError:
            pass
        for row in ws.iter_rows(max_row=row_limit):
            for c in row:
                v = getattr(c, "value", None)
                if v is None:
                    continue
                fmt = c.number_format or ""
                fmt = "" if fmt == "General" else fmt
                s = _classify_py(v, fmt, epoch1904, getattr(c, "data_type", None))
                if s is not None:
                    out[(ws.title, c.row, c.column)] = s
    wb.close()
    return out


def _classify_py(v, fmt: str, epoch1904: bool, data_type=None) -> Src | None:
    if isinstance(v, bool):
        return Src("bool", v, fmt)
    if isinstance(v, dt.datetime):
        if v < dt.datetime(1899, 12, 30):  # a negative serial: Excel shows ####, no date
            return Src("number", (v - dt.datetime(1899, 12, 30)).total_seconds() / 86400.0, fmt)
        if v == dt.datetime(1899, 12, 30) and not is_time_only_format(fmt):
            return Src("date", v.date(), fmt, 0)  # a date-formatted 0
        v = _round_second(v)
        if v.date() in (dt.date(1899, 12, 30), dt.date(1899, 12, 31)) and is_time_only_format(fmt):
            return Src("time", v.time(), fmt)
        if v.time() == dt.time(0):
            return Src("date", v.date(), fmt)
        return Src("datetime", v, fmt)
    if isinstance(v, dt.date):
        return Src("date", v, fmt)
    if isinstance(v, dt.time):
        return Src("time", v.replace(microsecond=0), fmt)
    if isinstance(v, dt.timedelta):
        return Src("time", (dt.datetime(2000, 1, 1) + v).time(), fmt)
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v != v:
            return None
        if is_date_format(fmt) and 0 <= v < 2958466:
            d = _date_cell(float(v), fmt, epoch1904)
            if d is not None:
                return d
        return Src("number", v, fmt)
    if isinstance(v, str):
        if "_x" in v:  # OOXML escapes openpyxl leaves in place: _x000D_ is a carriage return
            v = re.sub(r"_x([0-9A-Fa-f]{4})_", lambda m: chr(int(m.group(1), 16)), v)
        if data_type == "e":
            return Src("error", v, fmt)
        if v.strip() == "":
            return None
        return Src("text", v, fmt)
    return Src("text", str(v), fmt)


def read_xls(path: Path, row_limit: int) -> dict:
    import xlrd
    try:
        book = xlrd.open_workbook(str(path), formatting_info=True, on_demand=True)
        fmt_ok = True
    except Exception:
        book = xlrd.open_workbook(str(path), on_demand=True)
        fmt_ok = False
    out: dict = {}
    for si in range(book.nsheets):
        sh = book.sheet_by_index(si)
        for r in range(min(sh.nrows, row_limit)):
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                t, v = cell.ctype, cell.value
                if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                fmt = ""
                if fmt_ok:
                    try:
                        xf = book.xf_list[cell.xf_index]
                        fmt = book.format_map[xf.format_key].format_str or ""
                    except Exception:
                        fmt = ""
                fmt = "" if fmt == "General" else fmt
                key = (sh.name, r + 1, c + 1)
                if t == xlrd.XL_CELL_TEXT:
                    if v.strip():
                        out[key] = Src("text", v, fmt)
                elif t == xlrd.XL_CELL_NUMBER:
                    if is_date_format(fmt) and 0 <= v < 2958466:
                        d = _date_cell(v, fmt, book.datemode == 1)
                        out[key] = d or Src("number", v, fmt)
                    else:
                        out[key] = Src("number", v, fmt)
                elif t == xlrd.XL_CELL_DATE:
                    d = _date_cell(v, fmt or "yyyy-mm-dd", book.datemode == 1)
                    if d is not None:
                        if not is_date_format(fmt) and d.kind == "time":
                            d = Src("time", d.value, fmt, v)
                        out[key] = d
                elif t == xlrd.XL_CELL_BOOLEAN:
                    out[key] = Src("bool", bool(v), fmt)
                elif t == xlrd.XL_CELL_ERROR:
                    out[key] = Src("error", xlrd.error_text_from_code.get(v, "#ERR"), fmt)
        book.unload_sheet(si)
    book.release_resources()
    return out


def read_ods(path: Path, row_limit: int) -> dict:
    from odf.opendocument import load
    from odf import table, teletype, text
    doc = load(str(path))
    out: dict = {}
    for tb in doc.spreadsheet.getElementsByType(table.Table):
        name = tb.getAttribute("name")
        r = 0
        for row in _iter_rows(tb, table):
            rrep = int(row.getAttribute("numberrowsrepeated") or 1)
            cells = []
            c = 0
            for cell in row.childNodes:
                if cell.qname[1] not in ("table-cell", "covered-table-cell"):
                    continue
                crep = int(cell.getAttribute("numbercolumnsrepeated") or 1)
                s = _ods_cell(cell, teletype, text) if cell.qname[1] == "table-cell" else None
                if s is not None:
                    for k in range(min(crep, 1024)):
                        cells.append((c + k, s))
                c += crep
            if cells:
                for k in range(min(rrep, 1024)):
                    if r + k >= row_limit:
                        break
                    for cc, s in cells:
                        out[(name, r + k + 1, cc + 1)] = s
            r += rrep
            if r >= row_limit:
                break
    return out


def _iter_rows(tb, table):
    for node in tb.childNodes:
        q = node.qname[1]
        if q == "table-row":
            yield node
        elif q in ("table-row-group", "table-header-rows", "table-rows"):
            yield from _iter_rows(node, table)


_DUR = re.compile(r"^-?P(?:(\d+)D)?T?(?:(\d+)H)?(?:(\d+)M)?(?:([\d.]+)S)?$")


def _ods_cell(cell, teletype, text) -> Src | None:
    vt = cell.getAttribute("valuetype")
    if vt in ("float", "currency"):
        return Src("number", _num(cell.getAttribute("value")), "")
    if vt == "percentage":
        return Src("number", _num(cell.getAttribute("value")), "0%")
    if vt == "date":
        v = cell.getAttribute("datevalue") or ""
        try:
            d = dt.datetime.fromisoformat(v)
        except ValueError:
            return Src("text", v, "")
        d = _round_second(d)
        return Src("date", d.date(), "yyyy-mm-dd") if d.time() == dt.time(0) else Src("datetime", d, "yyyy-mm-dd hh:mm:ss")
    if vt == "time":
        m = _DUR.match(cell.getAttribute("timevalue") or "")
        if m:
            days, h, mi, s = (float(x or 0) for x in m.groups())
            t = dt.datetime(2000, 1, 1) + dt.timedelta(hours=h, minutes=mi, seconds=round(s))
            return Src("time", t.time(), "hh:mm:ss")
        return None
    if vt == "boolean":
        return Src("bool", (cell.getAttribute("booleanvalue") or "").lower() == "true", "")
    ps = cell.getElementsByType(text.P)
    if not ps:
        return None
    s = "\n".join(teletype.extractText(p) for p in ps)
    if not s.strip():
        return None
    if s.startswith("#") and s.upper() in ("#N/A", "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!"):
        return Src("error", s, "")
    return Src("text", s, "")


def _num(s):
    try:
        f = float(s)
    except (TypeError, ValueError):
        return 0
    return int(f) if f.is_integer() and abs(f) < 2**53 else f


def read_csv(path: Path, row_limit: int) -> dict:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            txt = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    lines = txt.splitlines()[:50]
    best, score = ",", -1
    for d in (",", ";", "\t", "|"):
        counts = [ln.count(d) for ln in lines if ln.strip()]
        if not counts or max(counts) == 0:
            continue
        common = max(set(counts), key=counts.count)
        s = counts.count(common) * (1 if common else 0)
        if s > score:
            best, score = d, s
    out: dict = {}
    for r, row in enumerate(csv.reader(io.StringIO(txt), delimiter=best)):
        if r >= row_limit:
            break
        for c, v in enumerate(row):
            if v.strip():
                out[("", r + 1, c + 1)] = Src("text", v, "")
    return out
