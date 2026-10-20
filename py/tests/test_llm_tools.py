"""Workbook tools on the docs sample, driven by a scripted model (no provider SDK, no network)."""

import json

import openpyxl
import pytest

import hedgehog
from hedgehog.llm import ChatClient, ChatResult, ToolCall, Usage, ask
from hedgehog.llm.tools import WorkbookTools


class Scripted(ChatClient):
    provider = "scripted"

    def __init__(self, replies):
        super().__init__(model="script-1", sleep=lambda s: None)
        self.replies = list(replies)

    def _chat(self, messages, **kw):
        r = self.replies.pop(0)
        if isinstance(r, str):
            return ChatResult(text=r, tool_calls=[], stop_reason="end_turn", usage=Usage(1, 1, calls=1), provider=self.provider)
        return ChatResult(text="", tool_calls=r, stop_reason="tool_use", usage=Usage(1, 1, calls=1), provider=self.provider)


@pytest.fixture
def quickstart_sales(tmp_path):
    """`sales.xlsx` from docs/quickstart.md: merged two-row header, group rows, subtotals, text numbers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Sales report 2024"
    ws.append([])
    ws.append(["Product", "Revenue", None, "Units", "Margin"])
    ws.append([None, "Q1", "Q2", None, None])
    for rng in ("B3:C3", "A3:A4", "D3:D4", "E3:E4"):
        ws.merge_cells(rng)
    for row in [
        ["Region: North"],
        ["Widgets", 1200.5, 1350, "1,200", "12%"],
        ["Gadgets", 300, 410, "350", "8.5%"],
        ["Subtotal", 1500.5, 1760, None, None],
        ["Region: South"],
        ["Widgets", 980, 1150, "n/a", "10%"],
        ["Gizmos", 2045, 1720, "2,100", "15%"],
        ["Subtotal", 3025, 2870, None, None],
        ["Grand total", 4600, 4630, None, None],
    ]:
        ws.append(row)
    ws.append([])
    ws.append(["Source: internal ERP export, amounts in USD"])
    path = tmp_path / "sales.xlsx"
    wb.save(path)
    return str(path)


def test_south_q2_revenue_tool_path(quickstart_sales):
    """The documented gpt-4o-mini run answered 0 after reading the empty Sales!C9 next to
    "Region: South". The tools return the right data on the path a model should take: the label,
    the block below it, the verified subtotal 2870 in C12, and the merged "Revenue" header over C."""
    llm = Scripted([
        [ToolCall("1", "find_text", {"query": "South"}), ToolCall("2", "find_text", {"query": "Q2"})],
        [ToolCall("3", "get_cell", {"sheet": "Sales", "address": "C9"}),
         ToolCall("4", "get_range", {"sheet": "Sales", "range": "A9:C12"}),
         ToolCall("5", "get_cell", {"sheet": "Sales", "address": "C3"}),
         ToolCall("6", "check_totals", {"sheet": "Sales"})],
        "South Q2 revenue is the South subtotal, 2870 (Sales!C12).\nANSWER: 2870",
    ])
    a = ask(quickstart_sales, "What is the Q2 revenue of the South region?", llm)
    results = {}
    for m in a.messages:
        for r in m.tool_results:
            assert not r.is_error, r.content
            results[r.tool_call_id] = json.loads(r.content)
    assert results["1"] == [{"cell": "Sales!A9", "value": "Region: South"}]
    assert results["2"] == [{"cell": "Sales!C4", "value": "Q2"}]
    # The cell next to the group label is empty (null), never 0.
    assert results["3"] == {"cell": "Sales!C9", "value": None, "type": "empty"}
    assert results["4"]["rows"] == [["Region: South", None, None], ["Widgets", 980, 1150], ["Gizmos", 2045, 1720], ["Subtotal", 3025, 2870]]
    # C3 is covered by the merged B3:C3 "Revenue" header.
    assert results["5"]["merged"] == "Sales!B3:C3" and results["5"]["merged_value"] == "Revenue"
    totals = results["6"]
    c12 = next(i for i in totals["items"] if i["cell"] == "Sales!C12")
    assert (c12["value"], c12["status"], c12["sources"], c12["label"]) == (2870, "ok", ["C10:C11"], "Subtotal")
    assert totals["items"][0]["cell"] == "Sales!B13" and totals["items"][0]["status"] == "mismatch"
    assert (totals["found"], totals["ok"], totals["mismatch"], totals["unknown"]) == (6, 5, 1, 0)
    assert a.final == "2870" and a.citations.ok and "Sales!C12" in a.cited_cells


def test_check_totals_agrees_with_manifest(tmp_path):
    """check_totals and the normalize manifest's Totals use the same statuses: an average in a
    total row is `unknown` (with the recomputed sum as `recomputed`, not `expected`), not a mismatch."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Q1", "Q2", "Share"])
    for row in [("North", 10, 20, 0.5), ("South", 5, 7, 0.25), ("East", 1, 2, 0.25)]:
        ws.append(row)
    ws.append(["Total", 16, 30, 1 / 3])  # Q2 is wrong (29); Share is an average, not a sum
    path = tmp_path / "t.xlsx"
    wb.save(path)
    manifest = hedgehog.normalize(str(path), str(tmp_path / "out")).manifest
    by_cell = {f"{m['sheet']}!{m['cell']}": m["status"] for m in manifest["totals"]["items"]}
    tools = json.loads(WorkbookTools(hedgehog.open(str(path))).run("check_totals", {})[0])
    assert {i["cell"]: i["status"] for i in tools["items"]} == by_cell
    assert by_cell == {"Sales!B5": "ok", "Sales!C5": "mismatch", "Sales!D5": "unknown"}
    assert (tools["ok"], tools["mismatch"], tools["unknown"]) == (1, 1, 1)
    assert tools["mismatch"] == manifest["totals"]["mismatch"]
    d5 = next(i for i in tools["items"] if i["cell"] == "Sales!D5")
    assert "expected" not in d5 and d5["recomputed"] == 1.0 and d5["reason"] == "matches_average"
