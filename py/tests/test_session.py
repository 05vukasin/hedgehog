import json
import threading

import pytest

import hedgehog


@pytest.fixture
def report(tmp_path):
    """The sales report of the quickstart: title, two header rows, group rows, subtotals, a wrong
    grand total and a source note."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Sales report 2024"
    ws.append([])
    ws.append(["Product", "Revenue", None, "Units"])
    ws.append([None, "Q1", "Q2", None])
    for rng in ("B3:C3", "A3:A4", "D3:D4"):
        ws.merge_cells(rng)
    for row in [
        ["Region: North"],
        ["Widgets", 1200.5, 1350, "1,200"],
        ["Gadgets", 300, 410, "350"],
        ["Subtotal", 1500.5, 1760, None],
        ["Region: South"],
        ["Widgets", 980, 1150, "n/a"],
        ["Gizmos", 2045, 1720, "2,100"],
        ["Subtotal", 3025, 2870, None],
        ["Grand total", 4600, 4630, None],
    ]:
        ws.append(row)
    ws.append([])
    ws.append(["Source: internal ERP export"])
    wb.save(tmp_path / "sales.xlsx")
    return tmp_path


def test_catalog_and_query(report):
    s = hedgehog.Session(report)
    cat = s.open_file("sales.xlsx")
    assert cat.startswith("FILE sales.xlsx")
    assert '"Revenue / Q1"' in cat and "Region[group]" in cat
    out = s.query("sales.xlsx", 'SELECT SUM("Revenue / Q1") FROM T1 WHERE Region = \'North\'')
    assert out.splitlines()[1] == "1500.5"
    allrows = s.query("sales.xlsx", "SELECT COUNT(*) FROM T1_all WHERE _row_type <> 'data'")
    assert allrows.splitlines()[1] == "3"
    assert s.open_files() == [("f1", "sales.xlsx")]


def test_other_tools(report):
    s = hedgehog.Session(report)
    assert "Widgets" in s.describe_table("sales.xlsx", "T1")
    assert "1200.5" in s.cells("sales.xlsx", "Sales", "B6")
    assert "mismatch" in s.totals("sales.xlsx", status="mismatch")
    assert "sales.xlsx" in s.list_files()


def test_question_and_call(report):
    s = hedgehog.Session(report)
    assert s.call("open_file", {"path": "sales.xlsx", "question": "Q1 revenue"}).startswith("FILE")
    out = s.call("query", {"path": "sales.xlsx", "sql": 'SELECT MAX("Revenue / Q2") FROM T1'})
    assert out.splitlines()[1] == "1720"


def test_errors_are_value_errors(report, tmp_path_factory):
    s = hedgehog.Session(report)
    with pytest.raises(ValueError, match="no such column"):
        s.query("sales.xlsx", "SELECT nope FROM T1")
    with pytest.raises(ValueError):
        s.query("sales.xlsx", "DELETE FROM T1")
    outside = tmp_path_factory.mktemp("outside") / "x.csv"
    outside.write_text("a,b\n1,2\n")
    with pytest.raises(ValueError):
        s.open_file(str(outside))
    s.allow(outside)
    assert s.open_file(str(outside)).startswith("FILE")
    with pytest.raises(ValueError):
        hedgehog.Session(report / "missing")


def test_limits(report):
    s = hedgehog.Session(report, max_cells=10, cache=False)
    with pytest.raises(ValueError):
        s.open_file("sales.xlsx")


def test_cache_dir(report, tmp_path_factory):
    d = tmp_path_factory.mktemp("cache")
    cat = hedgehog.Session(report, cache_dir=d).open_file("sales.xlsx")
    assert any(d.iterdir())
    assert hedgehog.Session(report, cache_dir=d).open_file("sales.xlsx").splitlines()[:3] == cat.splitlines()[:3]


def test_tool_specs():
    specs = hedgehog.tool_specs()
    names = [t["name"] for t in specs]
    assert {"open_file", "query", "describe_table", "cells", "totals", "list_files"} <= set(names)
    q = next(t for t in specs if t["name"] == "query")
    assert q["inputSchema"]["type"] == "object" and "sql" in q["inputSchema"]["properties"]
    json.dumps(specs)


def test_threads_share_a_session(report):
    s = hedgehog.Session(report)
    s.open_file("sales.xlsx")
    out = []
    ts = [threading.Thread(target=lambda: out.append(s.query("sales.xlsx", "SELECT COUNT(*) FROM T1"))) for _ in range(4)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    assert [o.splitlines()[1] for o in out] == ["4"] * 4
